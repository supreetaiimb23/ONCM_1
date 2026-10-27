"""Excel (.xlsx) export built from live database data, scoped by role."""
from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.models import (AccessRequest, AuditLog, ChangeLog, ChecklistItem, Document, DocumentReview, Notification,
                        Organization, Task, User, Workstream)
from app.services import STAGE_GROUP, by_id, doc_code, latest_maker_ids, visible_docs_clause

def _aw(x):
    return x if x is None or x.tzinfo else x.replace(tzinfo=timezone.utc)


HEAD = PatternFill("solid", fgColor="12302B")


def _sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append(["" if v is None else (v.replace(tzinfo=None) if isinstance(v, datetime) and v.tzinfo else v) for v in r])
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200]]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(max(width, 12), 55)


def build_workbook(db: Session, user: User, org_id=None, actor_type=None, stage=None, doc_status=None,
                   reviewer_id=None, date_from=None, date_to=None) -> tuple[bytes, str]:
    oq = db.query(Organization)
    if user.role == "secondary_admin":
        oq = oq.filter(Organization.assigned_admin_id == user.id)
    if org_id:
        oq = oq.filter(Organization.id == org_id)
    if actor_type:
        oq = oq.filter(Organization.actor_type == actor_type)
    if stage:
        oq = oq.filter(Organization.stage == stage)
    orgs = oq.order_by(Organization.name).all()
    oids = [o.id for o in orgs]
    admins = {u.id: u.name for u in db.query(User).filter(User.role != "participant").all()}
    onames = {o.id: o.name for o in orgs}

    dq = db.query(Document).filter(Document.organization_id.in_(oids)) if oids else db.query(Document).filter(False)
    if doc_status:
        dq = dq.filter(Document.status == doc_status)
    if reviewer_id:
        dq = dq.filter(Document.assigned_reviewer_id == reviewer_id)
    if date_from:
        dq = dq.filter(Document.created_at >= date_from)
    if date_to:
        dq = dq.filter(Document.created_at <= date_to)
    docs = dq.order_by(Document.id).all()
    dids = [d.id for d in docs]
    dmap = {d.id: d for d in docs}

    cl_done, cl_total = {}, {}
    if oids:
        for oid, done, n in (db.query(ChecklistItem.organization_id, ChecklistItem.done, func.count())
                             .filter(ChecklistItem.organization_id.in_(oids))
                             .group_by(ChecklistItem.organization_id, ChecklistItem.done)):
            cl_total[oid] = cl_total.get(oid, 0) + n
            if done:
                cl_done[oid] = cl_done.get(oid, 0) + n

    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Participant Tracker",
           ["Organization", "Actor type", "Stage", "Waiting on", "ONCM contact", "Contact email", "Account status",
            "Checklist done", "Last updated"],
           [[o.name, o.actor_type, o.stage, o.waiting_on, admins.get(o.assigned_admin_id), o.contact_email, o.status,
             "{}/{}".format(cl_done.get(o.id, 0), cl_total.get(o.id, 0)), o.updated_at] for o in orgs])
    _sheet(wb, "Document Register",
           ["Document ID", "Code", "Organization", "Folder", "Title", "Type", "Version", "Status", "Uploaded", "Reviewer", "Due date"],
           [[d.id, doc_code(d), onames.get(d.organization_id), d.category, d.title, d.doc_type, d.current_version, d.status,
             d.created_at, admins.get(d.assigned_reviewer_id), d.due_date] for d in docs])
    # ---- knowledge repository (respects each admin's category permissions)
    aids = None if user.role == "master_admin" else [o.id for o in db.query(Organization).filter(Organization.assigned_admin_id == user.id)]
    wsn = {w.id: w.name for w in db.query(Workstream)}
    rq = db.query(Document)
    vis = visible_docs_clause(user, aids)
    if vis is not None:
        rq = rq.filter(vis)
    repo = [d for d in rq.order_by(Document.id).all()
            if (not date_from or _aw(d.created_at) >= _aw(date_from)) and (not date_to or _aw(d.created_at) <= _aw(date_to))]
    makers = latest_maker_ids(db, repo)
    repo_orgs = by_id(db, Organization, (d.organization_id for d in repo))
    _sheet(wb, "Repository Register",
           ["Code", "Title", "Category", "Access level", "Stage", "Status", "Workstream", "Maker", "Checker", "Approver",
            "Classification confirmed", "Organization", "Updated"],
           [[doc_code(d), d.title, d.repo_category, d.access_level, STAGE_GROUP.get(d.status), d.status, wsn.get(d.workstream_id),
             admins.get(makers[d.id]), admins.get(d.assigned_reviewer_id), admins.get(d.approver_id),
             d.classification_confirmed_at, repo_orgs[d.organization_id].name if d.organization_id else "", d.updated_at] for d in repo])
    visible_ids = {d.id for d in repo}
    cl = [c for c in db.query(ChangeLog).order_by(ChangeLog.id.desc()).limit(2000)
          if all(i in visible_ids for i in (c.document_id, c.supersedes_document_id) if i)]
    cl_docs = by_id(db, Document, [i for c in cl for i in (c.document_id, c.supersedes_document_id)])
    _sheet(wb, "Change Log", ["When", "Title", "Decision", "Document", "Supersedes", "Workstream", "Recorded by"],
           [[c.created_at, c.title, c.decision, doc_code(cl_docs[c.document_id]) if c.document_id else "",
             doc_code(cl_docs[c.supersedes_document_id]) if c.supersedes_document_id else "",
             wsn.get(c.workstream_id), admins.get(c.recorded_by)] for c in cl])
    ws_docs = dict(db.query(Document.workstream_id, func.count()).filter(Document.workstream_id != None)  # noqa: E711
                   .group_by(Document.workstream_id).all())
    ws_tasks = dict(db.query(Task.workstream_id, func.count()).filter(Task.workstream_id != None, Task.status != "done")  # noqa: E711
                    .group_by(Task.workstream_id).all())
    _sheet(wb, "Workstreams", ["Workstream", "Owner", "Status", "Documents", "Open tasks"],
           [[w.name, admins.get(w.owner_id), w.status, ws_docs.get(w.id, 0), ws_tasks.get(w.id, 0)]
            for w in db.query(Workstream).order_by(Workstream.name)])
    revs = (db.query(DocumentReview).filter(DocumentReview.document_id.in_(dids)).order_by(DocumentReview.id).all()
            if dids else [])
    _sheet(wb, "Approval History",
           ["Document ID", "Title", "Organization", "Version", "Decision", "Reviewer", "Recorded reason", "Decided at"],
           [[r.document_id, dmap[r.document_id].title, onames.get(dmap[r.document_id].organization_id), r.version,
             r.decision, admins.get(r.reviewer_id), r.internal_comment or r.feedback, r.created_at] for r in revs])
    tq = db.query(Task).filter(Task.status != "done")
    if user.role == "secondary_admin":
        tq = tq.filter((Task.owner_id == user.id) | (Task.organization_id.in_(oids)))
    tasks = [t for t in tq.order_by(Task.due_date).all() if not org_id or t.organization_id == org_id]
    pend = [["Task", t.title, onames.get(t.organization_id, "-"), admins.get(t.owner_id), t.due_date, t.priority, t.status]
            for t in tasks]
    pend += [["Missing document", c.title, onames.get(c.organization_id), "Participant", None, "", "outstanding"]
             for c in db.query(ChecklistItem).filter(ChecklistItem.organization_id.in_(oids), ChecklistItem.done == False).all()] if oids else []  # noqa: E712
    _sheet(wb, "Pending Tasks", ["Kind", "Item", "Organization", "Owner", "Deadline", "Priority", "Status"], pend)
    if user.role == "master_admin" or user.can_export:
        ns = db.query(Notification).order_by(Notification.id.desc()).limit(2000).all()
        _sheet(wb, "Automation Report", ["Event", "Recipient", "Status", "Attempts", "Last error", "Created"],
               [[n.event, n.recipient, n.status, n.attempts, n.last_error, n.created_at] for n in ns])
    if user.role == "master_admin" or user.can_share:
        ar = db.query(AccessRequest).order_by(AccessRequest.id.desc()).limit(2000).all()
        ar_docs = by_id(db, Document, (r.document_id for r in ar))
        _sheet(wb, "Access Requests", ["Requested", "Document", "Requester", "Email", "Organization", "Source", "Status", "Decided", "Valid until"],
               [[r.created_at, doc_code(ar_docs[r.document_id]), r.requester_name, r.requester_email, r.requester_org,
                 r.source, r.status, r.decided_at, r.expires_at] for r in ar])
    if user.role == "master_admin":
        logs = db.query(AuditLog).order_by(AuditLog.id.desc()).limit(5000).all()
        _sheet(wb, "Audit Report", ["When", "Actor", "Action", "Resource", "Resource ID", "Detail"],
               [[a.created_at, a.actor_email, a.action, a.resource, a.resource_id, str(a.detail)] for a in logs])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"ONCM_Workbench_Report_{datetime.now(timezone.utc):%Y-%m-%d}.xlsx"
